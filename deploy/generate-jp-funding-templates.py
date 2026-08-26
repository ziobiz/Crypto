"""Generate KR/US/JP/CH/TH Excel templates for JP subsidiary funding docs."""
from __future__ import annotations

import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
LOCALES = ("KR", "US", "JP", "CH", "TH")
FONTS = {
    "KR": "Malgun Gothic",
    "JP": "Meiryo UI",
    "US": "Calibri",
    "CH": "Microsoft YaHei",
    "TH": "Tahoma",
}

THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
NAVY = PatternFill("solid", fgColor="1E3A5F")
GOLD = PatternFill("solid", fgColor="C9A227")
YELLOW = PatternFill("solid", fgColor="FFF3BF")
BLUE = PatternFill("solid", fgColor="DBEAFE")
GREEN = PatternFill("solid", fgColor="D1FAE5")
GRAY = PatternFill("solid", fgColor="F1F5F9")
WRAP = Alignment(wrap_text=True, vertical="center")

S = {
    "KR": {
        "cover_title": "일본 자회사 6개월 거래 예정 보고서 (TINPASS / 일본 계좌 입금)",
        "cover_body": (
            "본 양식은 일본 자회사가 일본 내 거래처로부터 자금을 수취한 뒤, "
            "TINPASS 일본 계좌에 입금하고, 그 대가로 자회사 대표에게 USDT를 지급하는 흐름을 "
            "국세청·AML 관점에서 사전에 설명하기 위한 보고서입니다.\n\n"
            "작성 주체: 일본 자회사\n대상기간: 8월부터 6개월 (8월~익년 1월)\n"
            "통화: JPY (일일 규모를 노란 칸에 입력하면 주간·월간은 수식으로 자동 합산)\n\n"
            "작성 순서\n1) 01_volume — 월별 일일 규모만 입력 (주간=일일×7, 월간=일일×해당월 일수)\n"
            "2) 02_daily — 필요 시 일자별 실적/예정\n3) 03_nature — 자금 성격·상대방·계약\n"
            "4) 04_docs — 개인/법인 제출·보관 서류\n5) 작성본을 저장 후 TINPASS USDT 신청 시 업로드"
        ),
        "legend": "노란 칸 = 직접 입력 / 파란 칸 = 자동계산 / 초록 칸 = 6개월 합계",
        "vol_title": "6개월 거래 규모 예정 (일일 입력 → 주간·월간 자동합산)",
        "vol_heads": ["월", "해당월 일수", "일일 거래규모 (JPY) ※입력", "주간 거래규모 (JPY) = 일일×7", "월간 거래규모 (JPY) = 일일×일수", "비고 (예정 건수·상대방 수)", "자금 성격 코드", "작성자 확인"],
        "sum6": "6개월 합계",
        "sum_note": "월간 합계(E열)가 실제 일본 계좌 입금 예정 총액의 기준입니다.",
        "avg_daily": "일평균 일일규모",
        "avg_month": "월평균 월간규모",
        "co_name": "작성 법인명",
        "rep_name": "대표자명",
        "date": "작성일",
        "sign": "인감/서명",
        "vol_note": "주간 규모는 일일×7, 월간 규모는 일일×해당월 일수입니다. 주말·휴일 미거래는 02_daily에서 일자별로 입력하십시오.",
        "codes": '"A_매출대금,B_용역보수,C_로열티,D_대여금회수,E_증자/출자,F_기타"',
        "code_err": "코드 목록에서 선택",
        "code_title": "자금 성격",
        "chart": "월간 거래규모 (JPY)",
        "daily_title": "일자별 예정 규모 (선택) — 입력 시 월합계가 01_volume 월간과 대사됩니다",
        "daily_heads": ["일자", "요일", "일일 규모(JPY) ※입력", "주 시작일(월)", "해당 주 합계", "월"],
        "wd": ["월", "화", "수", "목", "금", "토", "일"],
        "month_sum": "월별 합계 (02_daily)",
        "sheet_month": "01_volume 월간",
        "diff": "차이",
        "no_daily": "(일별미입력)",
        "nat_title": "거래 예상 성격 — 일본 자회사가 수취·당사 일본계좌에 입금하는 자금의 성격",
        "nat_heads": ["코드", "자금 성격", "상대방 유형(개인/법인)", "상대방 업종·관계", "계약/근거 문서", "예상 비중(%)", "비고(원천·과세·원천징수 여부)"],
        "nat_rows": [
            ("A", "일본 내 매출대금 (상품·재화)", "법인 또는 개인사업자", "", "매매계약·청구서·납품서", "", "소비세 과세 여부 기재"),
            ("B", "용역·컨설팅 보수", "법인 또는 개인", "", "업무위탁계약·청구서", "", "원천징수 대상 여부"),
            ("C", "로열티·라이선스", "법인", "", "라이선스 계약", "", "사용료 원천 검토"),
            ("D", "대여금 회수 / 그룹 대출", "관계회사", "", "금전소비대차·상환스케줄", "", "이자 약정·이전가격"),
            ("E", "증자·출자·자본납입", "주주", "", "주주총회의사록·납입증명", "", "자본거래 증빙"),
            ("F", "기타 (구체적 기재 필수)", "", "", "", "", "성격 불명이면 입금 불가"),
        ],
        "share_sum": "비중 합계(%) — 100이어야 함",
        "flow_h": "TINPASS 일본 계좌 입금 시 자금 흐름 요약",
        "flow": "① 일본 거래처 → 일본 자회사 계좌 (수취 시 상대방 KYC·계약·청구서 보관)\n② 일본 자회사 → TINPASS 지정 일본 은행계좌 (입금자명=자회사 명의와 일치)\n③ TINPASS → 일본 자회사 대표 지정 지갑으로 USDT 지급\n각 단계의 계약·송금명세·영수증은 일본 국세청 조사 시 즉시 제출할 수 있도록 보관합니다.",
        "docs_title": "제출·보관 서류 목록 (개인 / 법인 구분) — 자회사 내부 보관 및 TINPASS 제출",
        "docs_heads": ["구분", "단계", "서류명 (JP)", "서류명 (로컬)", "개인", "법인", "보관기간(권고)", "TINPASS 신청 시 업로드"],
        "req": "필수",
        "na": "해당없음",
        "keep": "업로드 불요·보관",
        "docs_items": [
            ("신원", "자회사 KYC", "本人確認書類（免許証/在留/パスポート）", "본인확인서류", "필수", "대표자 필수", "거래종료+7년", "사본"),
            ("신원", "자회사 KYC", "登記事項証明書・定款", "등기사항증명·정관", "해당없음", "필수", "최신본 상시", "사본"),
            ("신원", "자회사 KYC", "実質的支配者申告", "실질적 지배자 신고", "해당없음", "필수", "변경 시 갱신", "사본"),
            ("수취", "상대방→자회사", "取引契約書", "거래 계약서", "필수", "필수", "거래종료+7년", "요약 또는 전문"),
            ("수취", "상대방→자회사", "請求書・納品書", "청구서·납품서", "필수", "필수", "7년", "해당 건"),
            ("수취", "상대방→자회사", "相手方本人確認/登記", "상대방 본인확인/등기", "개인상대 시", "법인상대 시", "7년", "요청 시"),
            ("수취", "상대방→자회사", "振込依頼人名が確認できる入金明細", "입금명세(의뢰인명)", "필수", "필수", "7년", "권고"),
            ("이체", "자회사→TINPASS", "取締役決定・送金稟議", "이사회/송금 품의", "해당없음", "필수", "7년", "필수"),
            ("이체", "자회사→TINPASS", "送金依頼書・振込明細書", "송금의뢰·이체명세", "필수", "필수", "7년", "입금 후 별도"),
            ("이체", "자회사→TINPASS", "本 6か月予定報告書", "본 6개월 예정 보고서", "필수", "필수", "당해+6년", "필수"),
            ("세무", "국세청 대비", "総勘定元帳・現金出納", "총계정원장·출납장", "간이장부", "필수", "7년(법인税法)", "업로드 불요·보관"),
            ("세무", "국세청 대비", "消費税・法人税申告書控", "소비세·법인세 신고서 부본", "해당 시", "필수", "7년", "업로드 불요·보관"),
            ("세무", "국세청 대비", "移転価格・国外関連者書類", "이전가격·국외특수관계인", "해당 시", "해당 시", "7년", "요청 시"),
            ("암호자산", "USDT 수령", "代表者ウォレット所有証明", "대표 지갑 소유 증명", "필수", "필수", "거래종료+7년", "주소 등록본"),
            ("암호자산", "USDT 수령", "USDT受領後の帳簿記載", "수령 후 장부 기장", "필수", "필수", "7년", "업로드 불요·보관"),
        ],
        "docs_note": "개인: 자회사 대표가 개인사업 형태이거나 개인으로부터 수취하는 경우 양쪽 열을 확인하십시오.\n법인: 등기·실질적지배자·이사회 결의 없이 TINPASS 일본 계좌 입금을 진행하지 마십시오.\n보관: 일본 법인세법상 장부 7년. 국제조세·이전가격 해당 시 별도 문서 유지.",
        "app_title": "TINPASS USDT 매입 신청 — 일본 자회사 입금 근거 신청서",
        "col_item": "항목",
        "col_fill": "기재",
        "col_hint": "안내",
        "app_fields": [
            ("신청일", ""),
            ("신청 법인(일본 자회사) 정식명칭", ""),
            ("법인번호 / 등기번호", ""),
            ("대표자 성명", ""),
            ("대표자 생년월일", ""),
            ("본사 소재지", ""),
            ("담당자 성명·전화·이메일", ""),
            ("TINPASS 고객유형 (개인/법인)", ""),
            ("입금 예정 은행 — TINPASS 일본 계좌", "지정 계좌와 동일 명의로만 입금"),
            ("입금 예정 금액 (JPY)", ""),
            ("대응 USDT 수량 (예정)", ""),
            ("수령 지갑 주소 (대표자)", ""),
            ("네트워크 (TRC20/ERC20 등)", ""),
            ("자금 성격 코드 (A~F)", "보고서 03_nature와 일치"),
            ("일본 내 원 자금 제공자 (상호/성명)", ""),
            ("원 자금 제공자 유형 (개인/법인)", ""),
            ("원 자금 제공자와의 계약명·일자", ""),
            ("본 신청이 6개월 예정 범위 내인지", "예 / 아니오"),
            ("대표자 확인 (서명)", "허위 신고 시 거래 거절·동결될 수 있음"),
        ],
        "app_note": "본 신청서는 작성 후 PDF 또는 엑셀로 저장하고, 작성 완료한 6개월 거래 예정 보고서와 함께 TINPASS USDT 신청 화면에서 업로드합니다. 입금 영수증은 이체 완료 후 별도 제출합니다.",
        "chk_title": "TINPASS 신청 시 업로드 목록 (고객 유형별)",
        "chk_heads": ["No", "파일", "개인 고객", "법인 고객", "형식", "업로드 목적 코드", "체크"],
        "after": "입금 후",
        "chk_rows": [
            "6개월 거래 예정 보고서 (본 패키지 양식 작성본)",
            "본 신청서 01_apply 작성본",
            "원 자금 계약서 (자회사↔일본 거래처)",
            "원 자금 청구서·입금명세 (거래처→자회사)",
            "대표 본인확인서류",
            "등기사항증명·정관",
            "실질적 지배자 신고서",
            "송금 품의·이사회 결정 (자회사→TINPASS)",
            "TINPASS 일본 계좌 입금 영수증",
        ],
        "min_ind": "개인 고객 최소 세트: 1+2+3+4+5 (입금 후 9)",
        "min_corp": "법인 고객 최소 세트: 1+2+3+4+5+6+7+8 (입금 후 9)",
        "keep_note": "업로드하지 않고 자회사만 보관: 장부, 세금신고서 부본, USDT 수령 후 기장 — 세무조사 시 제출",
        "keep_title": "일본 자회사 보관 서류 (TINPASS 미업로드 · 조사 시 제출)",
        "keep_heads": ["분류", "서류", "근거(개요)", "보관 기한", "담당"],
        "keep_rows": [
            ("회계", "총계정원장, 보조원장, 현금출납장", "법인税法 장부서류", "7년", ""),
            ("회계", "입출금 통장·이체명세 원본", "자금 추적", "7년", ""),
            ("계약", "거래처 계약·청구·납품 원본", "익금산입·손금", "7년", ""),
            ("계약", "TINPASS 이용약관·신청 부본", "국외 관련 거래", "7년", ""),
            ("세무", "법인세·소비세 신고서 및 첨부", "신고 의무", "7년", ""),
            ("세무", "원천징수 관련 조서", "지급조서", "7년", ""),
            ("국조", "국외특수관계인 서류, 이전가격 문서", "해당 시", "7년", ""),
            ("암호자산", "지갑 키 관리 기록, TXID 목록", "자산 평가·양도", "7년", ""),
            ("AML", "거래처 KYC 사본, 제재 스크리닝 기록", "범죄수익", "거래종료+5년 이상", ""),
        ],
    },
}

# Fill remaining locales below after KR block in code — see locale_extra()


def locale_extra() -> dict:
    return {
        "US": {
            "cover_title": "Japan subsidiary 6-month trading forecast (TINPASS / Japan bank deposit)",
            "cover_body": (
                "This form explains in advance, for tax and AML, how the Japan subsidiary receives funds from Japanese counterparties, "
                "deposits them to the TINPASS Japan account, and receives USDT for the representative.\n\n"
                "Prepared by: Japan subsidiary\nPeriod: 6 months from August (Aug–Jan)\n"
                "Currency: JPY (enter daily volume in yellow cells; weekly/monthly auto-sum)\n\n"
                "Steps\n1) 01_volume — enter daily amount per month (weekly=daily×7, monthly=daily×days)\n"
                "2) 02_daily — optional day-by-day\n3) 03_nature — nature of funds\n"
                "4) 04_docs — individual vs corporate documents\n5) Save and upload when applying for USDT on TINPASS"
            ),
            "legend": "Yellow = input / Blue = formula / Green = 6-month total",
            "vol_title": "6-month volume forecast (daily input → weekly/monthly auto-sum)",
            "vol_heads": ["Month", "Days in month", "Daily volume (JPY) *input", "Weekly (JPY) = daily×7", "Monthly (JPY) = daily×days", "Notes (deals / counterparties)", "Fund-nature code", "Preparer"],
            "sum6": "6-month total",
            "sum_note": "Column E monthly totals are the planned Japan-account deposit amount.",
            "avg_daily": "Average daily",
            "avg_month": "Average monthly",
            "co_name": "Company name",
            "rep_name": "Representative",
            "date": "Date",
            "sign": "Seal / signature",
            "vol_note": "Weekly = daily×7; monthly = daily×days in month. Use 02_daily for weekends/holidays with no trades.",
            "codes": '"A_sales,B_service,C_royalty,D_loan_repay,E_capital,F_other"',
            "code_err": "Choose from the list",
            "code_title": "Nature of funds",
            "chart": "Monthly volume (JPY)",
            "daily_title": "Optional daily plan — monthly totals reconcile to 01_volume",
            "daily_heads": ["Date", "Weekday", "Daily (JPY) *input", "Week start (Mon)", "Week total", "Month"],
            "wd": ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"],
            "month_sum": "Monthly totals (02_daily)",
            "sheet_month": "01_volume monthly",
            "diff": "Difference",
            "no_daily": "(no daily input)",
            "nat_title": "Expected nature of funds received by the Japan subsidiary and deposited to TINPASS",
            "nat_heads": ["Code", "Nature", "Counterparty (individual/corporate)", "Industry / relationship", "Contract / evidence", "Share (%)", "Notes (tax / withholding)"],
            "nat_rows": [
                ("A", "Domestic sales (goods)", "Corporate or sole proprietor", "", "Sale contract / invoice / delivery", "", "Consumption tax"),
                ("B", "Service / consulting fees", "Corporate or individual", "", "Service agreement / invoice", "", "Withholding"),
                ("C", "Royalty / license", "Corporate", "", "License agreement", "", "Royalty withholding"),
                ("D", "Loan repayment / group loan", "Related company", "", "Loan agreement / schedule", "", "Interest / transfer pricing"),
                ("E", "Capital contribution", "Shareholder", "", "Minutes / payment proof", "", "Capital evidence"),
                ("F", "Other (must specify)", "", "", "", "", "Unclear nature: do not deposit"),
            ],
            "share_sum": "Share total (%) — must be 100",
            "flow_h": "Fund flow into the TINPASS Japan account",
            "flow": "1) JP counterparty → Japan subsidiary (keep KYC, contract, invoice)\n2) Japan subsidiary → TINPASS Japan bank (payer name = subsidiary)\n3) TINPASS → USDT to the representative wallet\nKeep contracts, remittance slips and receipts for Japan NTA.",
            "docs_title": "Documents to submit and retain (individual vs corporate)",
            "docs_heads": ["Category", "Stage", "Name (JP)", "Name (local)", "Individual", "Corporate", "Retention", "Upload on TINPASS apply"],
            "req": "Required",
            "na": "N/A",
            "keep": "Do not upload — retain",
            "docs_items": [
                ("ID", "Sub KYC", "本人確認書類（免許証/在留/パスポート）", "ID document", "Required", "Rep. required", "7y after end", "Copy"),
                ("ID", "Sub KYC", "登記事項証明書・定款", "Registry / articles", "N/A", "Required", "Keep current", "Copy"),
                ("ID", "Sub KYC", "実質的支配者申告", "Beneficial owner", "N/A", "Required", "Update on change", "Copy"),
                ("Receive", "CP→sub", "取引契約書", "Trade contract", "Required", "Required", "7y after end", "Summary or full"),
                ("Receive", "CP→sub", "請求書・納品書", "Invoice / delivery", "Required", "Required", "7 years", "Related deal"),
                ("Receive", "CP→sub", "相手方本人確認/登記", "Counterparty ID/registry", "If individual CP", "If corporate CP", "7 years", "On request"),
                ("Receive", "CP→sub", "振込依頼人名が確認できる入金明細", "Inbound advice (payer name)", "Required", "Required", "7 years", "Recommended"),
                ("Send", "sub→TINPASS", "取締役決定・送金稟議", "Board / remittance approval", "N/A", "Required", "7 years", "Required"),
                ("Send", "sub→TINPASS", "送金依頼書・振込明細書", "Remittance slip", "Required", "Required", "7 years", "After deposit"),
                ("Send", "sub→TINPASS", "本 6か月予定報告書", "This 6-month forecast", "Required", "Required", "6 years", "Required"),
                ("Tax", "NTA", "総勘定元帳・現金出納", "Ledgers / cash book", "Simplified books", "Required", "7y Corp Tax Act", "Do not upload — retain"),
                ("Tax", "NTA", "消費税・法人税申告書控", "Tax return copies", "If applicable", "Required", "7 years", "Do not upload — retain"),
                ("Tax", "NTA", "移転価格・国外関連者書類", "TP / foreign related-party", "If applicable", "If applicable", "7 years", "On request"),
                ("Crypto", "USDT", "代表者ウォレット所有証明", "Wallet ownership", "Required", "Required", "7y after end", "Registered address"),
                ("Crypto", "USDT", "USDT受領後の帳簿記載", "Book after receipt", "Required", "Required", "7 years", "Do not upload — retain"),
            ],
            "docs_note": "Individual: check both columns if the representative is a sole trader or funds come from individuals.\nCorporate: do not deposit to TINPASS without registry, beneficial owner and board approval.\nRetention: 7 years under Japan Corporate Tax Act.",
            "app_title": "TINPASS USDT application — Japan subsidiary funding evidence",
            "col_item": "Item",
            "col_fill": "Entry",
            "col_hint": "Guidance",
            "app_fields": [
                ("Application date", ""),
                ("Japan subsidiary legal name", ""),
                ("Corporate / registry number", ""),
                ("Representative name", ""),
                ("Representative date of birth", ""),
                ("Head office address", ""),
                ("Contact name / phone / email", ""),
                ("TINPASS customer type (individual/corporate)", ""),
                ("Planned bank — TINPASS Japan account", "Same name as the designated account only"),
                ("Planned amount (JPY)", ""),
                ("Expected USDT amount", ""),
                ("Receiving wallet (representative)", ""),
                ("Network (TRC20/ERC20 etc.)", ""),
                ("Fund-nature code (A–F)", "Must match 03_nature"),
                ("Original fund provider in Japan (name)", ""),
                ("Provider type (individual/corporate)", ""),
                ("Contract name and date with provider", ""),
                ("Within the 6-month forecast?", "Yes / No"),
                ("Representative confirmation (sign)", "False statements may lead to rejection or freeze"),
            ],
            "app_note": "Save as PDF or Excel and upload with the completed 6-month forecast on the TINPASS USDT application. Deposit receipts are submitted after the transfer.",
            "chk_title": "Files to upload on TINPASS apply (by customer type)",
            "chk_heads": ["No", "File", "Individual", "Corporate", "Format", "Purpose code", "Check"],
            "after": "After deposit",
            "chk_rows": [
                "Completed 6-month forecast (this pack)",
                "Completed application (01_apply)",
                "Origin contract (subsidiary ↔ JP counterparty)",
                "Origin invoice / inbound advice (CP → subsidiary)",
                "Representative ID",
                "Registry / articles",
                "Beneficial owner declaration",
                "Remittance approval / board decision (sub → TINPASS)",
                "TINPASS Japan account deposit receipt",
            ],
            "min_ind": "Individual minimum: 1+2+3+4+5 (then 9 after deposit)",
            "min_corp": "Corporate minimum: 1+2+3+4+5+6+7+8 (then 9 after deposit)",
            "keep_note": "Keep at subsidiary only: books, tax return copies, USDT booking — produce on tax audit",
            "keep_title": "Documents retained by the Japan subsidiary (not uploaded to TINPASS)",
            "keep_heads": ["Class", "Document", "Basis", "Retention", "Owner"],
            "keep_rows": [
                ("Accounting", "General / sub ledgers, cash book", "Corp. Tax Act books", "7 years", ""),
                ("Accounting", "Bank books / remittance originals", "Fund trail", "7 years", ""),
                ("Contract", "CP contracts / invoices / delivery originals", "Income / expense", "7 years", ""),
                ("Contract", "TINPASS terms / application copy", "Cross-border", "7 years", ""),
                ("Tax", "Corp. / consumption tax returns", "Filing duty", "7 years", ""),
                ("Tax", "Withholding statements", "Payment records", "7 years", ""),
                ("Intl tax", "Foreign related-party / TP files", "If applicable", "7 years", ""),
                ("Crypto", "Wallet key logs, TXID list", "Valuation / transfer", "7 years", ""),
                ("AML", "CP KYC copies, screening", "Proceeds of crime", "5y+ after end", ""),
            ],
        },
        "JP": {
            "cover_title": "日本子会社 6か月取引予定報告書（TINPASS / 日本口座入金）",
            "cover_body": (
                "本様式は、日本子会社が日本の取引先から資金を受領し、TINPASS日本口座へ入金した対価として代表者へUSDTを交付する流れを、税務・AMLの観点から事前に説明する報告書です。\n\n"
                "作成主体: 日本子会社\n対象期間: 8月から6か月（8月〜翌年1月）\n"
                "通貨: JPY（黄色セルに日次規模を入力すると週次・月次は数式で自動合計）\n\n"
                "手順\n1) 01_volume — 月別の日次のみ入力（週次=日次×7、月次=日次×日数）\n"
                "2) 02_daily — 必要に応じ日別\n3) 03_nature — 資金の性質\n"
                "4) 04_docs — 個人/法人の提出・保管書類\n5) 保存後、TINPASSのUSDT申請時にアップロード"
            ),
            "legend": "黄=入力 / 青=自動計算 / 緑=6か月合計",
            "vol_title": "6か月取引規模予定（日次入力→週次・月次自動合計）",
            "vol_heads": ["月", "当該月日数", "日次規模 (JPY) ※入力", "週次 (JPY)=日次×7", "月次 (JPY)=日次×日数", "備考（件数・相手先数）", "資金性質コード", "作成者確認"],
            "sum6": "6か月合計",
            "sum_note": "E列の月次合計が日本口座入金予定総額の基準です。",
            "avg_daily": "日次平均",
            "avg_month": "月次平均",
            "co_name": "作成法人名",
            "rep_name": "代表者名",
            "date": "作成日",
            "sign": "印鑑/署名",
            "vol_note": "週次は日次×7、月次は日次×当該月日数です。休業日は02_dailyで日別に入力してください。",
            "codes": '"A_売上代金,B_役務報酬,C_ロイヤリティ,D_貸付回収,E_増資/出資,F_その他"',
            "code_err": "リストから選択",
            "code_title": "資金の性質",
            "chart": "月次取引規模 (JPY)",
            "daily_title": "日別予定（任意）— 月合計は01_volumeと突合",
            "daily_heads": ["日付", "曜日", "日次(JPY) ※入力", "週開始(月)", "当該週合計", "月"],
            "wd": ["月", "火", "水", "木", "金", "土", "日"],
            "month_sum": "月別合計 (02_daily)",
            "sheet_month": "01_volume 月次",
            "diff": "差額",
            "no_daily": "(日別未入力)",
            "nat_title": "取引予定の性質 — 日本子会社が受領し当社日本口座へ入金する資金",
            "nat_heads": ["コード", "資金の性質", "相手方（個人/法人）", "業種・関係", "契約/根拠", "想定比率(%)", "備考（源泉・課税）"],
            "nat_rows": [
                ("A", "国内売上代金（商品）", "法人または個人事業主", "", "売買契約・請求・納品", "", "消費税"),
                ("B", "役務・コンサル報酬", "法人または個人", "", "業務委託契約・請求", "", "源泉徴収"),
                ("C", "ロイヤリティ・ライセンス", "法人", "", "ライセンス契約", "", "使用料源泉"),
                ("D", "貸付回収 / グループ貸付", "関係会社", "", "金銭消費貸借・返済表", "", "利息・移転価格"),
                ("E", "増資・出資", "株主", "", "株主総会議事録・払込証明", "", "資本取引"),
                ("F", "その他（具体記載必須）", "", "", "", "", "性質不明は入金不可"),
            ],
            "share_sum": "比率合計(%) — 100であること",
            "flow_h": "TINPASS日本口座入金時の資金フロー",
            "flow": "① 日本の取引先 → 日本子会社口座（受領時に相手方KYC・契約・請求を保管）\n② 日本子会社 → TINPASS指定日本口座（依頼人名=子会社名義）\n③ TINPASS → 代表者ウォレットへUSDT\n各段階の契約・送金明細・領収書は国税調査時に提出できるよう保管。",
            "docs_title": "提出・保管書類（個人/法人）— 子会社保管およびTINPASS提出",
            "docs_heads": ["区分", "段階", "書類名 (JP)", "書類名 (現地)", "個人", "法人", "保管期間", "TINPASS申請時アップロード"],
            "req": "必須",
            "na": "該当なし",
            "keep": "アップロード不要・保管",
            "docs_items": [
                ("身元", "子会社KYC", "本人確認書類（免許証/在留/パスポート）", "本人確認書類", "必須", "代表者必須", "取引終了+7年", "写し"),
                ("身元", "子会社KYC", "登記事項証明書・定款", "登記・定款", "該当なし", "必須", "最新を常時", "写し"),
                ("身元", "子会社KYC", "実質的支配者申告", "実質的支配者", "該当なし", "必須", "変更時更新", "写し"),
                ("受領", "相手→子会社", "取引契約書", "取引契約書", "必須", "必須", "取引終了+7年", "要約または全文"),
                ("受領", "相手→子会社", "請求書・納品書", "請求・納品", "必須", "必須", "7年", "当該件"),
                ("受領", "相手→子会社", "相手方本人確認/登記", "相手方本人確認/登記", "個人相手時", "法人相手時", "7年", "要請時"),
                ("受領", "相手→子会社", "振込依頼人名が確認できる入金明細", "入金明細（依頼人名）", "必須", "必須", "7年", "推奨"),
                ("送金", "子会社→TINPASS", "取締役決定・送金稟議", "取締役会/送金稟議", "該当なし", "必須", "7年", "必須"),
                ("送金", "子会社→TINPASS", "送金依頼書・振込明細書", "送金依頼・振込明細", "必須", "必須", "7年", "入金後別途"),
                ("送金", "子会社→TINPASS", "本 6か月予定報告書", "本6か月予定報告書", "必須", "必須", "当該+6年", "必須"),
                ("税務", "国税", "総勘定元帳・現金出納", "総勘定元帳・出納", "簡易帳簿", "必須", "7年(法人税法)", "アップロード不要・保管"),
                ("税務", "国税", "消費税・法人税申告書控", "消費税・法人税申告控", "該当時", "必須", "7年", "アップロード不要・保管"),
                ("税務", "国税", "移転価格・国外関連者書類", "移転価格・国外関連者", "該当時", "該当時", "7年", "要請時"),
                ("暗号資産", "USDT受領", "代表者ウォレット所有証明", "ウォレット所有証明", "必須", "必須", "取引終了+7年", "登録住所"),
                ("暗号資産", "USDT受領", "USDT受領後の帳簿記載", "受領後の記帳", "必須", "必須", "7年", "アップロード不要・保管"),
            ],
            "docs_note": "個人: 代表が個人事業、または個人から受領する場合は両列を確認。\n法人: 登記・実質的支配者・取締役会決議なしにTINPASS日本口座へ入金しない。\n保管: 法人税法上7年。国際課税・移転価格は別途文書。",
            "app_title": "TINPASS USDT購入申請 — 日本子会社入金根拠申請書",
            "col_item": "項目",
            "col_fill": "記載",
            "col_hint": "案内",
            "app_fields": [
                ("申請日", ""),
                ("申請法人（日本子会社）正式名称", ""),
                ("法人番号 / 登記番号", ""),
                ("代表者氏名", ""),
                ("代表者生年月日", ""),
                ("本店所在地", ""),
                ("担当者氏名・電話・メール", ""),
                ("TINPASS顧客類型（個人/法人）", ""),
                ("入金予定銀行 — TINPASS日本口座", "指定口座と同一名義のみ"),
                ("入金予定額 (JPY)", ""),
                ("対応USDT数量（予定）", ""),
                ("受領ウォレット（代表者）", ""),
                ("ネットワーク（TRC20/ERC20等）", ""),
                ("資金性質コード (A~F)", "03_natureと一致"),
                ("日本国内の原資提供者（商号/氏名）", ""),
                ("原資提供者の類型（個人/法人）", ""),
                ("原資提供者との契約名・日付", ""),
                ("本申請が6か月予定の範囲内か", "はい / いいえ"),
                ("代表者確認（署名）", "虚偽申告は拒否・凍結の可能性"),
            ],
            "app_note": "本申請書をPDFまたはExcelで保存し、記入済み6か月予定報告書とともにTINPASS USDT申請画面でアップロード。入金領収書は振込後に別途提出。",
            "chk_title": "TINPASS申請時アップロード一覧（顧客類型別）",
            "chk_heads": ["No", "ファイル", "個人顧客", "法人顧客", "形式", "目的コード", "チェック"],
            "after": "入金後",
            "chk_rows": [
                "6か月取引予定報告書（本パック記入済み）",
                "本申請書 01_apply 記入済み",
                "原資契約書（子会社↔日本取引先）",
                "原資請求・入金明細（取引先→子会社）",
                "代表者本人確認書類",
                "登記事項証明書・定款",
                "実質的支配者申告書",
                "送金稟議・取締役会決定（子会社→TINPASS）",
                "TINPASS日本口座入金領収書",
            ],
            "min_ind": "個人最小セット: 1+2+3+4+5（入金後9）",
            "min_corp": "法人最小セット: 1+2+3+4+5+6+7+8（入金後9）",
            "keep_note": "アップロードせず子会社保管: 帳簿、申告控、USDT受領記帳 — 税務調査時に提出",
            "keep_title": "日本子会社保管書類（TINPASS未アップロード・調査時提出）",
            "keep_heads": ["分類", "書類", "根拠", "保管期限", "担当"],
            "keep_rows": [
                ("会計", "総勘定元帳、補助元帳、現金出納帳", "法人税法帳簿", "7年", ""),
                ("会計", "入出金通帳・振込明細原本", "資金追跡", "7年", ""),
                ("契約", "取引先契約・請求・納品原本", "益金・損金", "7年", ""),
                ("契約", "TINPASS約款・申請控", "国外関連", "7年", ""),
                ("税務", "法人税・消費税申告及び添付", "申告義務", "7年", ""),
                ("税務", "源泉徴収関連調書", "支払調書", "7年", ""),
                ("国際", "国外関連者・移転価格文書", "該当時", "7年", ""),
                ("暗号資産", "ウォレット鍵管理、TXID一覧", "評価・譲渡", "7年", ""),
                ("AML", "取引先KYC写し、スクリーニング", "犯罪収益", "取引終了+5年以上", ""),
            ],
        },
        "CH": {
            "cover_title": "日本子公司6个月交易预估报告（TINPASS / 日本账户入金）",
            "cover_body": (
                "本表格用于事先说明：日本子公司从日本交易对手收款后，向 TINPASS 日本账户入金，并据此向代表支付 USDT，以便税务与 AML 审查。\n\n"
                "填写主体：日本子公司\n期间：自8月起6个月（8月–次年1月）\n"
                "币种：JPY（在黄色单元格填写日均规模，周/月由公式自动合计）\n\n"
                "步骤\n1) 01_volume — 仅填每月日均（周=日×7，月=日×当月天数）\n"
                "2) 02_daily — 可选按日\n3) 03_nature — 资金性质\n"
                "4) 04_docs — 个人/法人文件\n5) 保存后在 TINPASS USDT 申请时上传"
            ),
            "legend": "黄=填写 / 蓝=自动计算 / 绿=6个月合计",
            "vol_title": "6个月交易规模预估（填写日均 → 周/月自动合计）",
            "vol_heads": ["月", "当月天数", "日均规模 (JPY) ※填写", "周规模 (JPY)=日×7", "月规模 (JPY)=日×天数", "备注（笔数/对手数）", "资金性质代码", "填写人确认"],
            "sum6": "6个月合计",
            "sum_note": "E列月合计为日本账户计划入金总额基准。",
            "avg_daily": "日均平均",
            "avg_month": "月均平均",
            "co_name": "填写法人名称",
            "rep_name": "代表人",
            "date": "填写日",
            "sign": "印章/签名",
            "vol_note": "周=日×7，月=日×当月天数。周末/假日无交易请在 02_daily 按日填写。",
            "codes": '"A_货款,B_服务费,C_特许权,D_还款,E_增资,F_其他"',
            "code_err": "请从列表选择",
            "code_title": "资金性质",
            "chart": "月交易规模 (JPY)",
            "daily_title": "按日预估（可选）— 月合计与 01_volume 核对",
            "daily_heads": ["日期", "星期", "日规模(JPY) ※填写", "周起始(一)", "该周合计", "月"],
            "wd": ["一", "二", "三", "四", "五", "六", "日"],
            "month_sum": "按月合计 (02_daily)",
            "sheet_month": "01_volume 月计",
            "diff": "差额",
            "no_daily": "(未按日填写)",
            "nat_title": "交易预估性质 — 子公司收取并汇入 TINPASS 日本账户的资金",
            "nat_heads": ["代码", "资金性质", "对手类型（个人/法人）", "行业/关系", "合同/依据", "预计占比(%)", "备注（税/预提）"],
            "nat_rows": [
                ("A", "日本国内货款（商品）", "法人或个体户", "", "买卖合同/发票/交货", "", "消费税"),
                ("B", "服务/咨询费", "法人或个人", "", "委托合同/发票", "", "预提税"),
                ("C", "特许权/许可", "法人", "", "许可合同", "", "使用费预提"),
                ("D", "借款回收/集团贷款", "关联公司", "", "借贷合同/还款表", "", "利息/转让定价"),
                ("E", "增资/出资", "股东", "", "股东会纪要/缴款证明", "", "资本交易"),
                ("F", "其他（必须具体说明）", "", "", "", "", "性质不明不得入金"),
            ],
            "share_sum": "占比合计(%) — 须为100",
            "flow_h": "汇入 TINPASS 日本账户的资金流",
            "flow": "① 日本交易对手 → 日本子公司账户（收款时保管 KYC、合同、发票）\n② 日本子公司 → TINPASS 指定日本账户（汇款人名=子公司）\n③ TINPASS → 向代表人钱包支付 USDT\n各阶段合同、汇款明细、收据须备日本国税调查。",
            "docs_title": "提交与保管文件清单（个人/法人）",
            "docs_heads": ["分类", "阶段", "文件名 (JP)", "文件名 (本地)", "个人", "法人", "保管期限", "TINPASS 申请时上传"],
            "req": "必须",
            "na": "不适用",
            "keep": "无需上传·保管",
            "docs_items": [
                ("身份", "子公司KYC", "本人確認書類（免許証/在留/パスポート）", "身份证明", "必须", "代表人必须", "交易结束+7年", "副本"),
                ("身份", "子公司KYC", "登記事項証明書・定款", "登记/章程", "不适用", "必须", "保持最新", "副本"),
                ("身份", "子公司KYC", "実質的支配者申告", "实际控制人", "不适用", "必须", "变更时更新", "副本"),
                ("收款", "对手→子公司", "取引契約書", "交易合同", "必须", "必须", "交易结束+7年", "摘要或全文"),
                ("收款", "对手→子公司", "請求書・納品書", "发票/交货单", "必须", "必须", "7年", "相关交易"),
                ("收款", "对手→子公司", "相手方本人確認/登記", "对手身份/登记", "个人对手时", "法人对手时", "7年", "要求时"),
                ("收款", "对手→子公司", "振込依頼人名が確認できる入金明細", "入账明细（汇款人）", "必须", "必须", "7年", "建议"),
                ("汇出", "子公司→TINPASS", "取締役決定・送金稟議", "董事会/汇款审批", "不适用", "必须", "7年", "必须"),
                ("汇出", "子公司→TINPASS", "送金依頼書・振込明細書", "汇款申请/明细", "必须", "必须", "7年", "入金后另行"),
                ("汇出", "子公司→TINPASS", "本 6か月予定報告書", "本6个月预估报告", "必须", "必须", "当年+6年", "必须"),
                ("税务", "国税", "総勘定元帳・現金出納", "总账/出纳", "简易账", "必须", "7年(法人税)", "无需上传·保管"),
                ("税务", "国税", "消費税・法人税申告書控", "消费税/法人税申报副本", "适用时", "必须", "7年", "无需上传·保管"),
                ("税务", "国税", "移転価格・国外関連者書類", "转让定价/境外关联", "适用时", "适用时", "7年", "要求时"),
                ("加密资产", "USDT领取", "代表者ウォレット所有証明", "钱包所有权证明", "必须", "必须", "交易结束+7年", "登记地址"),
                ("加密资产", "USDT领取", "USDT受領後の帳簿記載", "领取后记账", "必须", "必须", "7年", "无需上传·保管"),
            ],
            "docs_note": "个人：代表为个体或从个人收款时请核对照栏。\n法人：无登记、实际控制人、董事会决议不得向 TINPASS 日本账户入金。\n保管：日本法人税账簿7年。",
            "app_title": "TINPASS USDT 申购 — 日本子公司入金依据申请书",
            "col_item": "项目",
            "col_fill": "填写",
            "col_hint": "说明",
            "app_fields": [
                ("申请日", ""),
                ("申请法人（日本子公司）正式名称", ""),
                ("法人号 / 登记号", ""),
                ("代表人姓名", ""),
                ("代表人生日", ""),
                ("总部地址", ""),
                ("经办人姓名·电话·邮箱", ""),
                ("TINPASS 客户类型（个人/法人）", ""),
                ("计划银行 — TINPASS 日本账户", "仅限与指定账户同名"),
                ("计划金额 (JPY)", ""),
                ("对应 USDT 数量（预计）", ""),
                ("收款钱包（代表人）", ""),
                ("网络（TRC20/ERC20 等）", ""),
                ("资金性质代码 (A~F)", "须与 03_nature 一致"),
                ("日本境内原资金提供方（名称）", ""),
                ("提供方类型（个人/法人）", ""),
                ("与提供方的合同名称及日期", ""),
                ("本申请是否在6个月预估范围内", "是 / 否"),
                ("代表人确认（签名）", "虚假申报可能导致拒绝或冻结"),
            ],
            "app_note": "填写后保存为 PDF 或 Excel，与已填 6 个月预估报告一并在 TINPASS USDT 申请页上传。入金收据在转账后另行提交。",
            "chk_title": "TINPASS 申请时上传清单（按客户类型）",
            "chk_heads": ["No", "文件", "个人客户", "法人客户", "格式", "用途代码", "勾选"],
            "after": "入金后",
            "chk_rows": [
                "6个月交易预估报告（本包已填）",
                "本申请书 01_apply 已填",
                "原资金合同（子公司↔日本对手）",
                "原资金发票/入账明细（对手→子公司）",
                "代表人身份证明",
                "登记事项证明·章程",
                "实际控制人申报",
                "汇款审批·董事会决定（子公司→TINPASS）",
                "TINPASS 日本账户入金收据",
            ],
            "min_ind": "个人最少：1+2+3+4+5（入金后9）",
            "min_corp": "法人最少：1+2+3+4+5+6+7+8（入金后9）",
            "keep_note": "不上传、由子公司保管：账簿、申报副本、USDT 入账 — 税务调查时提交",
            "keep_title": "日本子公司保管文件（不上传 TINPASS · 调查时提交）",
            "keep_heads": ["分类", "文件", "依据", "保管期限", "负责人"],
            "keep_rows": [
                ("会计", "总账、明细账、现金出纳", "法人税账簿", "7年", ""),
                ("会计", "进出账存折、汇款原件", "资金追踪", "7年", ""),
                ("合同", "对手合同、发票、交货原件", "益金/损金", "7年", ""),
                ("合同", "TINPASS 条款、申请副本", "境外相关", "7年", ""),
                ("税务", "法人税、消费税申报及附件", "申报义务", "7年", ""),
                ("税务", "预提相关书表", "支付记录", "7年", ""),
                ("国际", "境外关联方、转让定价文件", "适用时", "7年", ""),
                ("加密资产", "钱包密钥记录、TXID 列表", "估值/转让", "7年", ""),
                ("AML", "对手 KYC 副本、筛查记录", "犯罪收益", "交易结束+5年以上", ""),
            ],
        },
        "TH": {
            "cover_title": "รายงานคาดการณ์ธุรกรรม 6 เดือนของบริษัทย่อยญี่ปุ่น (TINPASS / บัญชีญี่ปุ่น)",
            "cover_body": (
                "แบบฟอร์มนี้ใช้อธิบายล่วงหน้าด้านภาษีและ AML ว่าบริษัทย่อยญี่ปุ่นรับเงินจากคู่ค้าในญี่ปุ่น แล้วโอนเข้าบัญชีญี่ปุ่นของ TINPASS และรับ USDT ให้ผู้แทน\n\n"
                "ผู้จัดทำ: บริษัทย่อยญี่ปุ่น\nช่วงเวลา: 6 เดือนจากสิงหาคม (ส.ค.–ม.ค.)\n"
                "สกุล: JPY (กรอกปริมาณรายวันในช่องเหลือง สัปดาห์/เดือนคำนวณอัตโนมัติ)\n\n"
                "ขั้นตอน\n1) 01_volume — กรอกเฉพาะรายวันต่อเดือน (สัปดาห์=วัน×7, เดือน=วัน×จำนวนวัน)\n"
                "2) 02_daily — รายวันตามต้องการ\n3) 03_nature — ลักษณะเงิน\n"
                "4) 04_docs — เอกสารบุคคล/นิติบุคคล\n5) บันทึกแล้วอัปโหลดตอนสมัคร USDT บน TINPASS"
            ),
            "legend": "เหลือง=กรอก / ฟ้า=สูตร / เขียว=รวม 6 เดือน",
            "vol_title": "คาดการณ์ปริมาณ 6 เดือน (กรอกรายวัน → สัปดาห์/เดือนอัตโนมัติ)",
            "vol_heads": ["เดือน", "วันในเดือน", "ปริมาณรายวัน (JPY) *กรอก", "รายสัปดาห์ (JPY)=วัน×7", "รายเดือน (JPY)=วัน×วันในเดือน", "หมายเหตุ (จำนวนรายการ/คู่ค้า)", "รหัสลักษณะเงิน", "ผู้จัดทำ"],
            "sum6": "รวม 6 เดือน",
            "sum_note": "ผลรวมรายเดือนคอลัมน์ E เป็นยอดที่จะโอนเข้าบัญชีญี่ปุ่น",
            "avg_daily": "เฉลี่ยรายวัน",
            "avg_month": "เฉลี่ยรายเดือน",
            "co_name": "ชื่อนิติบุคคล",
            "rep_name": "ผู้แทน",
            "date": "วันที่จัดทำ",
            "sign": "ตราประทับ/ลายเซ็น",
            "vol_note": "สัปดาห์=วัน×7 เดือน=วัน×วันในเดือน ใช้ 02_daily หากวันหยุดไม่มีธุรกรรม",
            "codes": '"A_ขาย,B_บริการ,C_สิทธิ,D_ชำระเงินกู้,E_เพิ่มทุน,F_อื่น"',
            "code_err": "เลือกจากรายการ",
            "code_title": "ลักษณะเงิน",
            "chart": "ปริมาณรายเดือน (JPY)",
            "daily_title": "แผนรายวัน (ไม่บังคับ) — รวมเดือนเทียบกับ 01_volume",
            "daily_heads": ["วันที่", "วัน", "รายวัน(JPY) *กรอก", "ต้นสัปดาห์(จ.)", "รวมสัปดาห์", "เดือน"],
            "wd": ["จ", "อ", "พ", "พฤ", "ศ", "ส", "อา"],
            "month_sum": "รวมรายเดือน (02_daily)",
            "sheet_month": "01_volume รายเดือน",
            "diff": "ผลต่าง",
            "no_daily": "(ไม่ได้กรอกรายวัน)",
            "nat_title": "ลักษณะเงินที่คาดหวัง — บริษัทย่อยรับแล้วโอนเข้าบัญชีญี่ปุ่น TINPASS",
            "nat_heads": ["รหัส", "ลักษณะเงิน", "ประเภทคู่ค้า (บุคคล/นิติ)", "อุตสาหกรรม/ความสัมพันธ์", "สัญญา/หลักฐาน", "สัดส่วน(%)", "หมายเหตุ (ภาษี/หัก ณ ที่จ่าย)"],
            "nat_rows": [
                ("A", "ยอดขายในญี่ปุ่น (สินค้า)", "นิติบุคคลหรือเจ้าของคนเดียว", "", "สัญญาซื้อขาย/ใบแจ้งหนี้/ส่งมอบ", "", "ภาษีบริโภค"),
                ("B", "ค่าบริการ/ที่ปรึกษา", "นิติบุคคลหรือบุคคล", "", "สัญญาจ้าง/ใบแจ้งหนี้", "", "หัก ณ ที่จ่าย"),
                ("C", "ค่าสิทธิ/ใบอนุญาต", "นิติบุคคล", "", "สัญญาใบอนุญาต", "", "หักค่าสิทธิ"),
                ("D", "ชำระเงินกู้ / กู้ในกลุ่ม", "บริษัทเกี่ยวข้อง", "", "สัญญากู้/ตารางชำระ", "", "ดอกเบี้ย/ราคาโอน"),
                ("E", "เพิ่มทุน/เงินลงทุน", "ผู้ถือหุ้น", "", "รายงานประชุม/หลักฐานชำระ", "", "หลักฐานทุน"),
                ("F", "อื่น (ต้องระบุ)", "", "", "", "", "ไม่ชัดเจนห้ามโอน"),
            ],
            "share_sum": "รวมสัดส่วน(%) — ต้องเป็น 100",
            "flow_h": "กระแสเงินเข้าบัญชีญี่ปุ่น TINPASS",
            "flow": "1) คู่ค้าญี่ปุ่น → บัญชีบริษัทย่อย (เก็บ KYC สัญญา ใบแจ้งหนี้)\n2) บริษัทย่อย → บัญชีธนาคารญี่ปุ่น TINPASS (ชื่อผู้โอน=บริษัทย่อย)\n3) TINPASS → USDT ไปวอลเล็ตผู้แทน\nเก็บสัญญา สลิป และใบเสร็จสำหรับสรรพากรญี่ปุ่น",
            "docs_title": "เอกสารยื่นและเก็บ (บุคคล / นิติบุคคล)",
            "docs_heads": ["หมวด", "ขั้น", "ชื่อเอกสาร (JP)", "ชื่อเอกสาร (ท้องถิ่น)", "บุคคล", "นิติบุคคล", "ระยะเก็บ", "อัปโหลดตอนสมัคร TINPASS"],
            "req": "จำเป็น",
            "na": "ไม่ใช้",
            "keep": "ไม่อัปโหลด·เก็บไว้",
            "docs_items": [
                ("ตัวตน", "KYC บริษัทย่อย", "本人確認書類（免許証/在留/パスポート）", "เอกสารยืนยันตัวตน", "จำเป็น", "ผู้แทนจำเป็น", "จบรายการ+7 ปี", "สำเนา"),
                ("ตัวตน", "KYC บริษัทย่อย", "登記事項証明書・定款", "ทะเบียน/ข้อบังคับ", "ไม่ใช้", "จำเป็น", "ฉบับล่าสุด", "สำเนา"),
                ("ตัวตน", "KYC บริษัทย่อย", "実質的支配者申告", "ผู้มีอำนาจควบคุมจริง", "ไม่ใช้", "จำเป็น", "อัปเดตเมื่อเปลี่ยน", "สำเนา"),
                ("รับเงิน", "คู่ค้า→บริษัทย่อย", "取引契約書", "สัญญาซื้อขาย", "จำเป็น", "จำเป็น", "จบรายการ+7 ปี", "สรุปหรือฉบับเต็ม"),
                ("รับเงิน", "คู่ค้า→บริษัทย่อย", "請求書・納品書", "ใบแจ้งหนี้/ส่งมอบ", "จำเป็น", "จำเป็น", "7 ปี", "รายการที่เกี่ยวข้อง"),
                ("รับเงิน", "คู่ค้า→บริษัทย่อย", "相手方本人確認/登記", "ตัวตน/ทะเบียนคู่ค้า", "ถ้าคู่ค้าบุคคล", "ถ้าคู่ค้านิติ", "7 ปี", "เมื่อขอ"),
                ("รับเงิน", "คู่ค้า→บริษัทย่อย", "振込依頼人名が確認できる入金明細", "รายการเข้า (ชื่อผู้โอน)", "จำเป็น", "จำเป็น", "7 ปี", "แนะนำ"),
                ("โอน", "บริษัทย่อย→TINPASS", "取締役決定・送金稟議", "มติกรรมการ/อนุมัติโอน", "ไม่ใช้", "จำเป็น", "7 ปี", "จำเป็น"),
                ("โอน", "บริษัทย่อย→TINPASS", "送金依頼書・振込明細書", "คำขอโอน/สลิป", "จำเป็น", "จำเป็น", "7 ปี", "หลังฝากแยก"),
                ("โอน", "บริษัทย่อย→TINPASS", "本 6か月予定報告書", "รายงานคาดการณ์ 6 เดือนนี้", "จำเป็น", "จำเป็น", "ปีนั้น+6 ปี", "จำเป็น"),
                ("ภาษี", "สรรพากร", "総勘定元帳・現金出納", "บัญชีแยกประเภท/เงินสด", "บัญชีอย่างง่าย", "จำเป็น", "7 ปี (ภาษีนิติ)", "ไม่อัปโหลด·เก็บไว้"),
                ("ภาษี", "สรรพากร", "消費税・法人税申告書控", "สำเนายื่นภาษี", "ถ้ามี", "จำเป็น", "7 ปี", "ไม่อัปโหลด·เก็บไว้"),
                ("ภาษี", "สรรพากร", "移転価格・国外関連者書類", "ราคาโอน/บุคคลต่างประเทศ", "ถ้ามี", "ถ้ามี", "7 ปี", "เมื่อขอ"),
                ("คริปโต", "รับ USDT", "代表者ウォレット所有証明", "หลักฐานถือวอลเล็ต", "จำเป็น", "จำเป็น", "จบรายการ+7 ปี", "ที่อยู่ที่ลงทะเบียน"),
                ("คริปโต", "รับ USDT", "USDT受領後の帳簿記載", "ลงบัญชีหลังรับ", "จำเป็น", "จำเป็น", "7 ปี", "ไม่อัปโหลด·เก็บไว้"),
            ],
            "docs_note": "บุคคล: ตรวจทั้งสองคอลัมน์หากผู้แทนเป็นเจ้าของคนเดียวหรือรับเงินจากบุคคล\nนิติบุคคล: ห้ามโอนเข้า TINPASS หากไม่มีทะเบียน ผู้มีอำนาจควบคุมจริง และมติกรรมการ\nเก็บ: บัญชี 7 ปีตามกฎหมายภาษีนิติบุคคลญี่ปุ่น",
            "app_title": "คำขอซื้อ USDT TINPASS — หลักฐานเงินฝากบริษัทย่อยญี่ปุ่น",
            "col_item": "รายการ",
            "col_fill": "กรอก",
            "col_hint": "คำแนะนำ",
            "app_fields": [
                ("วันที่ยื่น", ""),
                ("ชื่อนิติบุคคลบริษัทย่อยญี่ปุ่น", ""),
                ("เลขทะเบียนนิติบุคคล", ""),
                ("ชื่อผู้แทน", ""),
                ("วันเกิดผู้แทน", ""),
                ("ที่ตั้งสำนักงานใหญ่", ""),
                ("ผู้ติดต่อ ชื่อ/โทร/อีเมล", ""),
                ("ประเภทลูกค้า TINPASS (บุคคล/นิติ)", ""),
                ("ธนาคารที่จะโอน — บัญชีญี่ปุ่น TINPASS", "โอนเฉพาะชื่อเดียวกับบัญชีที่กำหนด"),
                ("ยอดที่จะโอน (JPY)", ""),
                ("จำนวน USDT ที่คาด", ""),
                ("วอลเล็ตรับ (ผู้แทน)", ""),
                ("เครือข่าย (TRC20/ERC20 ฯลฯ)", ""),
                ("รหัสลักษณะเงิน (A–F)", "ต้องตรงกับ 03_nature"),
                ("ผู้ให้เงินต้นในญี่ปุ่น (ชื่อ)", ""),
                ("ประเภทผู้ให้ (บุคคล/นิติ)", ""),
                ("ชื่อสัญญาและวันที่กับผู้ให้", ""),
                ("อยู่ในกรอบคาดการณ์ 6 เดือนหรือไม่", "ใช่ / ไม่"),
                ("ผู้แทนยืนยัน (ลายเซ็น)", "แจ้งเท็จอาจถูกปฏิเสธหรืออายัด"),
            ],
            "app_note": "บันทึกเป็น PDF หรือ Excel แล้วอัปโหลดพร้อมรายงานคาดการณ์ 6 เดือนที่กรอกแล้วในหน้าสมัคร USDT ของ TINPASS ใบเสร็จฝากส่งหลังโอน",
            "chk_title": "รายการอัปโหลดตอนสมัคร TINPASS (ตามประเภทลูกค้า)",
            "chk_heads": ["No", "ไฟล์", "ลูกค้าบุคคล", "ลูกค้านิติบุคคล", "รูปแบบ", "รหัสวัตถุประสงค์", "ติ๊ก"],
            "after": "หลังฝาก",
            "chk_rows": [
                "รายงานคาดการณ์ 6 เดือน (ชุดนี้ที่กรอกแล้ว)",
                "ใบสมัคร 01_apply ที่กรอกแล้ว",
                "สัญญาระเงินต้น (บริษัทย่อย↔คู่ค้าญี่ปุ่น)",
                "ใบแจ้งหนี้/รายการเข้าเงินต้น (คู่ค้า→บริษัทย่อย)",
                "เอกสารยืนยันตัวตนผู้แทน",
                "หนังสือรับรองทะเบียน/ข้อบังคับ",
                "แบบผู้มีอำนาจควบคุมจริง",
                "อนุมัติโอน/มติกรรมการ (บริษัทย่อย→TINPASS)",
                "ใบเสร็จฝากบัญชีญี่ปุ่น TINPASS",
            ],
            "min_ind": "ชุดขั้นต่ำบุคคล: 1+2+3+4+5 (หลังฝาก 9)",
            "min_corp": "ชุดขั้นต่ำนิติบุคคล: 1+2+3+4+5+6+7+8 (หลังฝาก 9)",
            "keep_note": "ไม่อัปโหลด เก็บที่บริษัทย่อย: บัญชี สำเนายื่นภาษี บันทึก USDT — ยื่นเมื่อถูกตรวจ",
            "keep_title": "เอกสารที่บริษัทย่อยญี่ปุ่นเก็บ (ไม่อัปโหลด TINPASS · ยื่นเมื่อถูกตรวจ)",
            "keep_heads": ["หมวด", "เอกสาร", "หลักเกณฑ์", "ระยะเก็บ", "ผู้รับผิดชอบ"],
            "keep_rows": [
                ("บัญชี", "บัญชีแยกประเภท บัญชีย่อย เงินสด", "บัญชีตามภาษีนิติ", "7 ปี", ""),
                ("บัญชี", "สมุดบัญชี/สลิปโอนต้นฉบับ", "ติดตามเงิน", "7 ปี", ""),
                ("สัญญา", "สัญญา คู่ค้า ใบแจ้งหนี้ ส่งมอบต้นฉบับ", "รายได้/ค่าใช้จ่าย", "7 ปี", ""),
                ("สัญญา", "ข้อกำหนด TINPASS สำเนาคำขอ", "รายการข้ามประเทศ", "7 ปี", ""),
                ("ภาษี", "แบบยื่นภาษีนิติ/บริโภคและเอกสารแนบ", "หน้าที่ยื่น", "7 ปี", ""),
                ("ภาษี", "เอกสารหัก ณ ที่จ่าย", "บันทึกการจ่าย", "7 ปี", ""),
                ("ภาษีระหว่างปท.", "เอกสารบุคคลต่างประเทศ/ราคาโอน", "ถ้ามี", "7 ปี", ""),
                ("คริปโต", "บันทึกกุญแจวอลเล็ต รายการ TXID", "ประเมิน/โอน", "7 ปี", ""),
                ("AML", "สำเนา KYC คู่ค้า การคัดกรอง", "เงินจากอาชญากรรม", "จบรายการ+อย่างน้อย 5 ปี", ""),
            ],
        },
    }


def t(loc: str) -> dict:
    data = dict(S["KR"])
    extra = locale_extra()
    if loc in extra:
        data.update(extra[loc])
    return data


def fonts(loc: str):
    name = FONTS.get(loc, "Calibri")
    return (
        Font(color="FFFFFF", bold=True, name=name, size=11),
        Font(bold=True, name=name, size=16, color="1E3A5F"),
        Font(bold=True, name=name, size=11),
        Font(name=name, size=10),
    )


def style_header_row(ws: Worksheet, row: int, cols: int, white: Font) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = NAVY
        cell.font = white
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN


def set_widths(ws: Worksheet, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def paint_input(cell, body: Font) -> None:
    cell.fill = YELLOW
    cell.border = THIN
    cell.font = body
    cell.alignment = Alignment(horizontal="right", vertical="center")


def formula_cell(cell, formula: str, body: Font, fill=BLUE) -> None:
    cell.value = formula
    cell.fill = fill
    cell.border = THIN
    cell.font = body
    cell.number_format = "#,##0"
    cell.alignment = Alignment(horizontal="right", vertical="center")


def make_forecast(loc: str) -> Workbook:
    s = t(loc)
    white, title, head, body = fonts(loc)
    wb = Workbook()
    cover = wb.active
    cover.title = "00_guide"
    cover["A1"] = s["cover_title"]
    cover["A1"].font = title
    cover.merge_cells("A1:F1")
    cover["A3"] = s["cover_body"]
    cover["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    cover.merge_cells("A3:F18")
    cover["A20"] = s["legend"]
    cover["A20"].fill = YELLOW
    set_widths(cover, [22, 22, 22, 22, 22, 40])
    cover.row_dimensions[1].height = 28
    cover.row_dimensions[3].height = 220

    vol = wb.create_sheet("01_volume")
    vol["A1"] = s["vol_title"]
    vol["A1"].font = title
    vol.merge_cells("A1:H1")
    for i, h in enumerate(s["vol_heads"], 1):
        vol.cell(3, i, h)
    style_header_row(vol, 3, 8, white)
    months = [("2026-08", 31), ("2026-09", 30), ("2026-10", 31), ("2026-11", 30), ("2026-12", 31), ("2027-01", 31)]
    for i, (ym, days) in enumerate(months):
        r = 4 + i
        vol.cell(r, 1, ym).font = head
        vol.cell(r, 1).border = THIN
        vol.cell(r, 2, days).border = THIN
        vol.cell(r, 2).alignment = Alignment(horizontal="center")
        paint_input(vol.cell(r, 3), body)
        vol.cell(r, 3).number_format = "#,##0"
        formula_cell(vol.cell(r, 4), f'=IF(C{r}="","",C{r}*7)', body)
        formula_cell(vol.cell(r, 5), f'=IF(C{r}="","",C{r}*B{r})', body)
        for c in (6, 7, 8):
            vol.cell(r, c).fill = YELLOW
            vol.cell(r, c).border = THIN
    vol.cell(10, 1, s["sum6"]).font = head
    vol.cell(10, 1).fill = GREEN
    vol.cell(10, 2, "=SUM(B4:B9)").fill = GREEN
    formula_cell(vol.cell(10, 3), "=SUM(C4:C9)", body, GREEN)
    formula_cell(vol.cell(10, 4), "=SUM(D4:D9)", body, GREEN)
    formula_cell(vol.cell(10, 5), "=SUM(E4:E9)", body, GREEN)
    vol.merge_cells("F10:H10")
    vol["F10"] = s["sum_note"]
    vol["F10"].fill = GREEN
    vol["A12"] = s["avg_daily"]
    formula_cell(vol["B12"], '=IF(COUNT(C4:C9)=0,"",AVERAGE(C4:C9))', body)
    vol["A13"] = s["avg_month"]
    formula_cell(vol["B13"], '=IF(COUNT(E4:E9)=0,"",AVERAGE(E4:E9))', body)
    vol["A15"] = s["co_name"]
    paint_input(vol["B15"], body)
    vol.merge_cells("B15:D15")
    vol["A16"] = s["rep_name"]
    paint_input(vol["B16"], body)
    vol["A17"] = s["date"]
    paint_input(vol["B17"], body)
    vol["C16"] = s["sign"]
    paint_input(vol["D16"], body)
    vol.merge_cells("D16:E16")
    vol["A19"] = s["vol_note"]
    vol.merge_cells("A19:H21")
    vol["A19"].alignment = WRAP
    set_widths(vol, [16, 14, 28, 28, 32, 32, 16, 14])
    vol.row_dimensions[3].height = 36
    vol.freeze_panes = "A4"
    dv = DataValidation(type="list", formula1=s["codes"], allow_blank=True)
    dv.error = s["code_err"]
    dv.errorTitle = s["code_title"]
    vol.add_data_validation(dv)
    dv.add("G4:G9")
    chart = BarChart()
    chart.title = s["chart"]
    chart.y_axis.title = "JPY"
    data = Reference(vol, min_col=5, min_row=3, max_row=9)
    cats = Reference(vol, min_col=1, min_row=4, max_row=9)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.y_axis.scaling.min = 0
    vol.add_chart(chart, "A23")

    daily = wb.create_sheet("02_daily")
    daily["A1"] = s["daily_title"]
    daily["A1"].font = title
    daily.merge_cells("A1:F1")
    for i, h in enumerate(s["daily_heads"], 1):
        daily.cell(3, i, h)
    style_header_row(daily, 3, 6, white)
    start = datetime.date(2026, 8, 1)
    end = datetime.date(2027, 1, 31)
    r = 4
    d = start
    weekdays = s["wd"]
    while d <= end:
        daily.cell(r, 1, d).number_format = "YYYY-MM-DD"
        daily.cell(r, 1).border = THIN
        daily.cell(r, 2, weekdays[d.weekday()]).border = THIN
        if d.weekday() >= 5:
            daily.cell(r, 2).fill = GRAY
        paint_input(daily.cell(r, 3), body)
        daily.cell(r, 3).number_format = "#,##0"
        daily.cell(r, 4, f"=A{r}-WEEKDAY(A{r},2)+1")
        daily.cell(r, 4).number_format = "YYYY-MM-DD"
        daily.cell(r, 4).fill = BLUE
        daily.cell(r, 4).border = THIN
        daily.cell(r, 5, f"=SUMIF($D$4:$D$187,D{r},$C$4:$C$187)")
        daily.cell(r, 5).number_format = "#,##0"
        daily.cell(r, 5).fill = BLUE
        daily.cell(r, 5).border = THIN
        daily.cell(r, 6, f'=TEXT(A{r},"YYYY-MM")').border = THIN
        r += 1
        d += datetime.timedelta(days=1)
    last = r - 1
    daily.cell(last + 2, 1, s["month_sum"]).font = head
    for i, (ym, _) in enumerate(months):
        rr = last + 3 + i
        daily.cell(rr, 1, ym)
        formula_cell(daily.cell(rr, 2), f"=SUMIF($F$4:$F${last},A{rr},$C$4:$C${last})", body)
        daily.cell(rr, 3, s["sheet_month"])
        formula_cell(daily.cell(rr, 4), f"=VLOOKUP(A{rr},'01_volume'!A4:E9,5,FALSE)", body)
        daily.cell(rr, 5, s["diff"])
        formula_cell(daily.cell(rr, 6), f'=IF(B{rr}=0,"{s["no_daily"]}",B{rr}-D{rr})', body, GREEN)
    set_widths(daily, [14, 10, 22, 16, 16, 14])
    daily.freeze_panes = "A4"
    daily.auto_filter.ref = f"A3:F{last}"

    nature = wb.create_sheet("03_nature")
    nature["A1"] = s["nat_title"]
    nature["A1"].font = title
    nature.merge_cells("A1:G1")
    for i, h in enumerate(s["nat_heads"], 1):
        nature.cell(3, i, h)
    style_header_row(nature, 3, 7, white)
    for i, row in enumerate(s["nat_rows"]):
        rr = 4 + i
        for c, v in enumerate(row, 1):
            nature.cell(rr, c, v).border = THIN
            nature.cell(rr, c).font = body
            if c in (3, 4, 6, 7):
                nature.cell(rr, c).fill = YELLOW
        nature.cell(rr, 6).number_format = "0.0"
    nature["A11"] = s["share_sum"]
    formula_cell(nature["B11"], "=SUM(F4:F9)", body, GREEN)
    nature["A13"] = s["flow_h"]
    nature["A13"].font = head
    nature["A14"] = s["flow"]
    nature.merge_cells("A14:G18")
    nature["A14"].alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(nature, [8, 32, 22, 22, 28, 14, 36])

    docs = wb.create_sheet("04_docs")
    docs["A1"] = s["docs_title"]
    docs["A1"].font = title
    docs.merge_cells("A1:H1")
    for i, h in enumerate(s["docs_heads"], 1):
        docs.cell(3, i, h)
    style_header_row(docs, 3, 8, white)
    req, keep = s["req"], s["keep"]
    for i, row in enumerate(s["docs_items"]):
        rr = 4 + i
        for c, v in enumerate(row, 1):
            cell = docs.cell(rr, c, v)
            cell.border = THIN
            cell.font = body
            cell.alignment = WRAP
            if v == req:
                cell.fill = GOLD
            elif keep in str(v):
                cell.fill = GRAY
        docs.row_dimensions[rr].height = 28
    docs["A21"] = s["docs_note"]
    docs.merge_cells("A21:H24")
    docs["A21"].alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(docs, [12, 18, 40, 28, 14, 14, 18, 22])
    return wb


def make_application(loc: str) -> Workbook:
    s = t(loc)
    white, title, head, body = fonts(loc)
    wb = Workbook()
    app = wb.active
    app.title = "01_apply"
    app["A1"] = s["app_title"]
    app["A1"].font = title
    app.merge_cells("A1:D1")
    app["A3"] = s["col_item"]
    app["B3"] = s["col_fill"]
    app["C3"] = s["col_hint"]
    style_header_row(app, 3, 3, white)
    for i, (label, hint) in enumerate(s["app_fields"]):
        r = 4 + i
        app.cell(r, 1, label).border = THIN
        app.cell(r, 1).fill = GRAY
        app.cell(r, 1).font = body
        paint_input(app.cell(r, 2), body)
        app.cell(r, 2).alignment = WRAP
        app.cell(r, 3, hint).border = THIN
        app.row_dimensions[r].height = 22
    app["A24"] = s["app_note"]
    app.merge_cells("A24:C26")
    app["A24"].alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(app, [36, 42, 48])

    chk = wb.create_sheet("02_upload")
    chk["A1"] = s["chk_title"]
    chk["A1"].font = title
    for i, h in enumerate(s["chk_heads"], 1):
        chk.cell(3, i, h)
    style_header_row(chk, 3, 7, white)
    req, na, after = s["req"], s["na"], s["after"]
    meta = [
        (req, req, "xlsx/PDF", "FUNDING_FORECAST_REPORT"),
        (req, req, "xlsx/PDF", "SOURCE_OF_FUNDS_DOC"),
        (req, req, "PDF", "SOURCE_OF_FUNDS_DOC"),
        (req, req, "PDF/img", "SOURCE_OF_FUNDS_DOC"),
        (req, req, "PDF/img", "SOURCE_OF_FUNDS_DOC"),
        (na, req, "PDF", "JP_TAX_SUPPORT_DOC"),
        (na, req, "PDF", "JP_TAX_SUPPORT_DOC"),
        (na, req, "PDF", "JP_TAX_SUPPORT_DOC"),
        (after, after, "PDF/img", "FIAT_DEPOSIT_RECEIPT"),
    ]
    for i, name in enumerate(s["chk_rows"]):
        rr = 4 + i
        ind, corp, fmt, code = meta[i]
        vals = (str(i + 1), name, ind, corp, fmt, code, "")
        for c, v in enumerate(vals, 1):
            cell = chk.cell(rr, c, v)
            cell.border = THIN
            if v == req:
                cell.fill = GOLD
            if c == 7:
                cell.fill = YELLOW
    chk["A15"] = s["min_ind"]
    chk["A16"] = s["min_corp"]
    chk["A17"] = s["keep_note"]
    set_widths(chk, [6, 48, 14, 14, 16, 28, 10])

    keep = wb.create_sheet("03_retain")
    keep["A1"] = s["keep_title"]
    keep["A1"].font = title
    for i, h in enumerate(s["keep_heads"], 1):
        keep.cell(3, i, h)
    style_header_row(keep, 3, 5, white)
    for i, row in enumerate(s["keep_rows"]):
        rr = 4 + i
        for c, v in enumerate(row, 1):
            keep.cell(rr, c, v).border = THIN
            if c == 5:
                keep.cell(rr, c).fill = YELLOW
    set_widths(keep, [12, 42, 28, 22, 16])
    return wb


def save_all() -> None:
    names = (
        "TINPASS_JP_6month_funding_forecast.xlsx",
        "TINPASS_JP_USDT_application_checklist.xlsx",
    )
    public = ROOT / "frontend" / "public" / "templates"
    docs = ROOT / "docs" / "templates"
    for loc in LOCALES:
        makers = (make_forecast(loc), make_application(loc))
        for base in (public / loc, docs / loc):
            base.mkdir(parents=True, exist_ok=True)
            for name, wb in zip(names, makers):
                path = base / name
                wb.save(path)
                print("wrote", path)
        if loc == "KR":
            public.mkdir(parents=True, exist_ok=True)
            docs.mkdir(parents=True, exist_ok=True)
            for name, wb in zip(names, makers):
                (public / name).parent.mkdir(parents=True, exist_ok=True)
                wb.save(public / name)
                wb.save(docs / name)
                print("wrote fallback", public / name)


if __name__ == "__main__":
    save_all()
